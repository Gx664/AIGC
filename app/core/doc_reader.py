import os


def extract_text(path):
    ext = os.path.splitext(path)[1].lower()
    if ext == ".txt":
        return _read_txt(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".pdf":
        return _read_pdf(path)
    raise ValueError("仅支持 PDF / DOCX / TXT 文件")


def _read_txt(path):
    for enc in ("utf-8", "gb18030"):
        try:
            with open(path, "r", encoding=enc) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _read_docx(path):
    from docx import Document

    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _read_pdf(path):
    from pypdf import PdfReader

    reader = PdfReader(path)
    parts = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts)


def split_paragraphs(text, min_len=20):
    raw = [line.strip() for line in text.splitlines()]
    paras = []
    buf = []
    for line in raw:
        if line:
            buf.append(line)
        else:
            if buf:
                paras.append("".join(buf))
                buf = []
    if buf:
        paras.append("".join(buf))
    merged = []
    for p in paras:
        if merged and len(p) < min_len:
            merged[-1] += p
        else:
            merged.append(p)
    return [p for p in merged if len(p) >= min_len]

# aigc-toolkit: file purpose marker
